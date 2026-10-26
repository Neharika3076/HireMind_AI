from pathlib import Path

import fitz
from docx import Document


def extract_from_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file.
    """

    text = []

    document = fitz.open(file_path)

    for page in document:
        page_text = page.get_text()

        if page_text:
            text.append(page_text)

    document.close()

    return "\n".join(text)


def extract_from_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file.
    """

    document = Document(file_path)

    text = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())

            if row_text:
                text.append(" | ".join(row_text))

    return "\n".join(text)


def extract_resume(file_path: str) -> str:
    """
    Detect the file type and extract resume text.
    """

    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return extract_from_pdf(file_path)

    if extension == ".docx":
        return extract_from_docx(file_path)

    raise ValueError(
        "Unsupported file format. Only PDF and DOCX are supported."
    )